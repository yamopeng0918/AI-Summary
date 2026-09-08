"""Verify the generated document and render every PDF slide for visual QA."""
from pathlib import Path
import sys, json, re, zipfile, hashlib
from pptx import Presentation
from PIL import Image, ImageOps, ImageDraw

root=Path(__file__).resolve().parent
sys.path.insert(0,str(root/'.build/python-deps'))
import pypdfium2 as pdfium

pptx=root/'output/AI-Digest-結案報告.pptx'
pdf=next((root/'.build/rendered').glob('*.pdf'))
prs=Presentation(pptx)
doc=pdfium.PdfDocument(pdf)
assert len(prs.slides)==len(doc)==9
report={'slides':9,'notes':0,'chartValues':[], 'outOfBounds':[], 'missingRenderedText':[], 'unexpectedTextOverlaps':[], 'rendered':[]}
for i,(sl,page) in enumerate(zip(prs.slides,doc),1):
    note=sl.notes_slide.notes_text_frame.text
    assert '[Sources]' in note and len(note)>200
    report['notes']+=1
    texts=[]
    visible=page.get_textpage().get_text_range()
    norm=lambda x:re.sub(r'\s+','',x)
    for sh in sl.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+20 or sh.top+sh.height>prs.slide_height+20:
            report['outOfBounds'].append([i,sh.name])
        if sh.has_chart:
            values=list(sh.chart.series[0].values)
            assert abs(values[0]-91.6666666667)<.001 and abs(values[1]-16.6666666667)<.001
            report['chartValues']=values
        if sh.has_text_frame and sh.text.strip():
            texts.append(sh)
            if norm(sh.text) not in norm(visible):
                report['missingRenderedText'].append({'slide':i,'text':sh.text})
    for a,sh in enumerate(texts):
        for other in texts[a+1:]:
            dx=min(sh.left+sh.width,other.left+other.width)-max(sh.left,other.left)
            dy=min(sh.top+sh.height,other.top+other.height)-max(sh.top,other.top)
            if dx>9144 and dy>9144:
                report['unexpectedTextOverlaps'].append({'slide':i,'a':sh.text,'b':other.text})
    target=root/f'.build/rendered/slide-{i:02}.png'
    page.render(scale=1.6).to_pil().save(target)
    report['rendered'].append(target.name)

with zipfile.ZipFile(pptx) as package:
    for name in package.namelist():
        if name.endswith('.xml'):
            raw=package.read(name).decode('utf-8')
            assert not re.search(r'(?:sk-[A-Za-z0-9_-]{20,}|ghp_[A-Za-z0-9]{20,}|AIza[A-Za-z0-9_-]{25,}|C:\\\\Users\\\\)',raw),name
report['sha256']=hashlib.sha256(pptx.read_bytes()).hexdigest()
report['sensitivePatternScan']='passed (known credential patterns; not proof against every possible secret)'
(root/'.build/verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=True,indent=2))
assert not report['outOfBounds']
assert not report['unexpectedTextOverlaps']
assert not report['missingRenderedText']
