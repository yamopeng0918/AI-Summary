"""Author the approved closure report; no product code or network calls."""
from pathlib import Path
import json
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / '.build/assets'
OUTPUT = ROOT / 'output/AI-Digest-結案報告.pptx'
FONT = 'Microsoft JhengHei'
BG, INK, ACCENT, MUTED, PALE = 'F5F0E6', '173D32', 'D75A43', '586B62', 'E8E4D8'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = 'AI Digest：結案報告'
prs.core_properties.subject = '三來源核心 MVP 成果與技術驗證'
prs.core_properties.author = '彭元懋'
prs.core_properties.keywords = 'AI Digest, 結案報告'
source = (ROOT / 'ai-digest-closure-slide-content.md').read_text(encoding='utf-8')
sections = re.split(r'^## \d+\. ', source, flags=re.M)[1:]
positions = []

def text(slide, value, x, y, w, h, size=20, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name, p.font.size = FONT, Pt(size)
        p.font.bold, p.font.color.rgb = bold, RGBColor.from_string(color)
        p.space_after = Pt(5)
        if align is not None:
            p.alignment = align
        for run in p.runs:
            rpr = run._r.get_or_add_rPr()
            ea = OxmlElement('a:ea')
            ea.set('typeface', FONT)
            rpr.append(ea)
    positions.append({'slide':len(prs.slides), 'text':value, 'x':x,'y':y,'w':w,'h':h,'fontPt':size})
    return shape

def line(slide,x1,y1,x2,y2,color=PALE,width=1):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=RGBColor.from_string(color)
    sh.line.width=Pt(width)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh

def box(slide,label,x,y,w,h,fill=PALE,size=18,color=INK):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill)
    sh.line.fill.background()
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    text(slide,label,x+.12,y+.16,w-.24,h-.2,size,color,True,PP_ALIGN.CENTER)

def picture(slide,name,x,y,w,h):
    from PIL import Image
    path=ASSETS/name
    with Image.open(path) as im:
        iw,ih=im.size
    scale=min(w/iw,h/ih)
    return slide.shapes.add_picture(str(path), Inches(x+(w-iw*scale)/2), Inches(y+(h-ih*scale)/2),width=Inches(iw*scale),height=Inches(ih*scale))

def new(title=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(INK if dark else BG)
    n=len(prs.slides)
    if title:
        text(s,title,.6,.43,12.1,.78,35,BG if dark else INK,True)
    text(s,'AI DIGEST  /  專案結案報告',.6,7.08,10,.24,10,PALE if dark else MUTED)
    text(s,f'{n:02d}',12,7.02,.65,.35,14,PALE if dark else MUTED,align=PP_ALIGN.RIGHT)
    note=re.split(r'### 講者備註（\d+ 秒）\s*',sections[n-1])[1].split('\n## 製作與驗證清單')[0].strip()
    # Resolve source references from the notes to repository paths, avoiding private absolute paths.
    note=note.replace('../../','').replace('../superpowers/','docs/superpowers/')
    if n in (1,3,5):
        note+='\n- 畫面來源：repository site/dist 的本機渲染，擷取日期 2026-09-05；對應公開網站 https://yamopeng0918.github.io/AI-Summary/'
    s.notes_slide.notes_text_frame.text=note
    return s

# 1 — Minimal cover with a real website image.
s=new()
text(s,'AI Digest',.6,1.02,6,1.12,60,bold=True)
text(s,'從公開網址到可驗證、\n可部署的 AI 知識摘要系統',.65,2.48,5.4,1.5,25)
text(s,'彭元懋',.65,5.67,5,.4,20,bold=True)
text(s,'專案結案報告  ·  2026.09',.65,6.16,5,.4,16,MUTED)
picture(s,'cover.png',6.45,1,6.3,5.3)

# 2 — Product problem and objective.
s=new('問題與專案目標')
text(s,'公開內容分散，\n閱讀後仍需要整理',.65,1.7,11.8,1.6,36,bold=True)
for x,label in [(0.65,'公開網址'),(4.95,'結構化繁中摘要'),(9.35,'搜尋與分類瀏覽')]:
    text(s,label,x,4,3.3,.7,23,bold=True)
text(s,'→',4,4,.7,.6,26,ACCENT)
text(s,'→',8.45,4,.7,.6,26,ACCENT)
text(s,'核心 MVP',.65,5.65,1.7,.4,18,ACCENT,True)
text(s,'一般網頁  ／  YouTube  ／  Bluesky 公開單篇貼文',2.45,5.65,10.2,.6,21)
text(s,'只處理可直接讀取的公開內容；不可讀時回報明確錯誤。',.65,6.4,12,.4,16,MUTED)

# 3 — Actual finished site and three concise results.
s=new('最終成果總覽')
overview=s.shapes.add_picture(str(ASSETS/'overview.png'), Inches(.65), Inches(1.75), width=Inches(8.15), height=Inches(4.584375))
overview.crop_top=.25
for y,num,title,body in [(1.75,'01','三種來源','網頁、影片、公開貼文'),(3.28,'02','結構化摘要','繁中摘要、重點與編輯觀點'),(4.81,'03','GitHub Pages','搜尋、分類、排序與詳情')]:
    text(s,num,9.03,y,.8,.5,24,ACCENT,True)
    text(s,title,9.03,y+.52,3.7,.48,23,bold=True)
    text(s,body,9.03,y+1.06,3.7,.4,16,MUTED)

# 4 — Only acquisition branches; all other stages share one pipeline.
s=new('三來源完整處理流程')
box(s,'公開網址',.65,2.45,1.45,.72)
box(s,'來源辨識',2.55,2.45,1.45,.72)
text(s,'→',2.13,2.57,.4,.5,22,ACCENT)
for y,label in [(1.5,'一般網頁｜正文'),(2.5,'YouTube｜字幕優先'),(3.5,'Bluesky｜公開單篇')]:
    box(s,label,4.8,y,3.9,.68,size=18)
    line(s,4.4,2.81,4.4,y+.34,MUTED)
    line(s,4.4,y+.34,4.8,y+.34,MUTED)
    line(s,8.7,y+.34,9.1,y+.34,MUTED)
    line(s,9.1,y+.34,9.1,2.84,MUTED)
line(s,4,2.81,4.4,2.81,MUTED)
text(s,'無可用字幕 → 音訊轉錄',4.95,3.2,3.8,.3,16,MUTED)
box(s,'取得文字',9.65,2.5,2.95,.68,fill=INK,color=BG)
line(s,9.1,2.84,9.65,2.84,MUTED)
text(s,'↓',10.82,3.38,.6,.5,24,ACCENT)
text(s,'共同處理',.65,4.68,2,.5,20,ACCENT,True)
text(s,'AI 摘要  →  分類  →  Schema 驗證  →  JSON 儲存',2.75,4.68,10,.6,22,bold=True)
text(s,'後續獨立操作',.65,5.77,2.1,.6,18,ACCENT,True)
text(s,'Astro 建置  →  GitHub Pages',2.75,5.75,9.8,.6,24)
text(s,'驗證通過才保存；新增摘要不會自動部署。',2.75,6.43,9.8,.4,16,MUTED)

# 5 — Three different real source outputs.
s=new('三來源成果展示')
for name,x,title,desc in [('web',.65,'一般網頁','從正文整理學習重點'),('youtube',4.95,'YouTube','字幕與音訊匯入同一格式'),('social',9.25,'Bluesky','保留貼文來源與作者')]:
    text(s,title,x,1.37,3.45,.48,25,bold=True)
    text(s,desc,x,1.95,3.45,.64,16,MUTED)
    picture(s,f'{name}.png',x,2.57,3.45,4.18)

# 6 — Responsibility boundaries, distinct from sequence slide.
s=new('系統架構與技術選擇')
text(s,'本機 Python CLI',.65,1.55,7.4,.7,30,bold=True)
text(s,'獨立解析器 → 摘要服務 → 分類器',.65,2.53,7.5,.6,24)
text(s,'來源故障隔離；各模組只負責自己的工作',.65,3.21,7.5,.6,18,MUTED)
text(s,'共同 Schema 與 JSON 儲存',.65,4.25,7.5,.7,27,bold=True)
text(s,'資料完整且通過驗證，才寫入獨立檔案',.65,5.07,7.5,.55,19,MUTED)
line(s,8.43,1.6,8.43,6.2,ACCENT,2)
text(s,'公開網站',9,1.55,3.7,.55,25,ACCENT,True)
text(s,'Astro\nGitHub Pages',9,2.47,3.7,1.25,29,bold=True)
text(s,'只讀取已驗證的\npublished 資料',9,4.13,3.7,1.1,21)
text(s,'前端不持有 API 金鑰',.65,6.13,11.9,.55,24,ACCENT,True)

# 7 — Native chart with editable data.
s=new('分類模型與量化證據')
text(s,'91.67%',.65,1.54,5.5,1.15,62,ACCENT,True)
text(s,'Accuracy｜33 / 36 筆正確',.68,2.83,5.35,.5,21)
text(s,'0.9179',.65,3.75,5.5,.9,43,bold=True)
text(s,'Macro F1',.68,4.72,5.1,.42,20,MUTED)
data=CategoryChartData(); data.categories=['分類模型','最大類基準']; data.add_series('Accuracy',[91.6666666667,16.6666666667])
chart=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(6.25), Inches(1.6), Inches(6.37), Inches(3.9), data).chart
chart.has_legend=False
chart.chart_style=10
chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=100; chart.value_axis.major_unit=25
chart.value_axis.tick_labels.font.name=FONT; chart.value_axis.tick_labels.font.size=Pt(14)
chart.category_axis.tick_labels.font.name=FONT; chart.category_axis.tick_labels.font.size=Pt(18)
chart.category_axis.reverse_order=True
chart.plots[0].has_data_labels=True
chart.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
chart.plots[0].data_labels.number_format='0.00"%"'
chart.plots[0].data_labels.font.name=FONT; chart.plots[0].data_labels.font.size=Pt(17)
for pt,col in zip(chart.series[0].points,[ACCENT,MUTED]):
    pt.format.fill.solid(); pt.format.fill.fore_color.rgb=RGBColor.from_string(col)
    pt.format.line.fill.background()
text(s,'Accuracy 高出基準 75 個百分點',6.25,5.6,6.4,.55,23,bold=True)
text(s,'180 筆人工審核資料，六類各 30 筆｜訓練 144／測試 36｜seed 42',.65,6.35,12,.5,17,MUTED)

# 8 — Open rows instead of UI cards.
s=new('跨來源整合挑戰與解法')
for y,no,challenge,solution,detail in [(1.65,'01','文字取得方式不同','獨立解析器與共同資料契約','單一平台失效，不破壞其他來源流程'),(3.32,'02','YouTube 無可用字幕','音訊轉錄與暫存清理','字幕優先；結束或失敗時處理本機與遠端清理'),(4.99,'03','公開部署的資料風險','驗證、掃描與部署 gate','不完整資料拒絕寫入，公開前檢查建置成果')]:
    text(s,no,.65,y,.8,.6,26,ACCENT,True)
    text(s,challenge,1.52,y,4.4,.6,23,bold=True)
    text(s,'→',6.05,y,.6,.5,24,ACCENT)
    text(s,solution,6.85,y,5.85,.6,23,bold=True)
    text(s,detail,6.85,y+.72,5.85,.65,17,MUTED)

# 9 — Evidence and conclusion; readable limits remain visible.
s=new('三來源核心 MVP 已完成',dark=True)
text(s,'715',.65,1.6,3.1,1.1,60,BG,True)
text(s,'Python passed · 2 skipped',.65,2.87,5.5,.5,18,PALE)
text(s,'67',6.8,1.6,3.1,1.1,60,BG,True)
text(s,'Vitest passed',6.8,2.87,5.5,.5,18,PALE)
text(s,'正式建置與部署驗證通過',.65,3.72,12,.6,27,BG,True)
text(s,'真實 push → matching workflow → 公開 smoke 通過',.65,4.45,12,.5,20,BG)
text(s,'學習成果｜資料契約、安全邊界、可重現評估與部署可靠性',.65,5.4,12,.7,20,BG)
text(s,'未實作：PDF／OCR／登入內容／網站後台',.65,6.24,12,.4,17,PALE)
text(s,'驗收依據：2026-09-03 進度紀錄',.65,6.74,9.5,.27,12,PALE)
link=text(s,'公開網站 ↗',10.4,6.64,2.2,.33,14,PALE)
link.click_action.hyperlink.address='https://yamopeng0918.github.io/AI-Summary/'

OUTPUT.parent.mkdir(exist_ok=True,parents=True)
prs.save(OUTPUT)
(ROOT/'.build/layout.json').write_text(json.dumps(positions,ensure_ascii=False,indent=2),encoding='utf-8')
assert len(prs.slides)==9
assert all('[Sources]' in sl.notes_slide.notes_text_frame.text for sl in prs.slides)
check=Presentation(OUTPUT)
assert len([sh for sh in check.slides[6].shapes if sh.has_chart])==1
for sl in check.slides:
    for sh in sl.shapes:
        assert sh.left>=0 and sh.top>=0
        assert sh.left+sh.width<=prs.slide_width+20 and sh.top+sh.height<=prs.slide_height+20
print('Saved nine-slide editable PPTX; notes, native chart, and canvas bounds verified.')
